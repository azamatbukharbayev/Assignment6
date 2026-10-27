#!/usr/bin/env python3
"""
Assignment 6: Comprehensive Academic Report Generator
Author: Azamat Bukharbayev (IT-2510)
Produces: Assignment6_AzamatBukharbayev.docx and Assignment6_AzamatBukharbayev.pdf
"""

import os
import sys
import subprocess
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

base_dir = "/Users/azamatbukharbayev/Documents/AITU/Web(Frontend)/Assignment6"
docx_path = os.path.join(base_dir, "Assignment6_AzamatBukharbayev.docx")
pdf_path = os.path.join(base_dir, "Assignment6_AzamatBukharbayev.pdf")

doc = docx.Document()

# Page Margins: 1 inch on all sides
for section in doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

# Colors
COLOR_PRIMARY = RGBColor(15, 23, 42)      # Deep slate (#0F172A)
COLOR_GOLD = RGBColor(217, 119, 6)        # Amber / Gold accent (#D97706)
COLOR_MUTED = RGBColor(100, 116, 139)     # Slate muted (#64748B)
COLOR_TEXT = RGBColor(30, 41, 59)         # Dark charcoal text (#1E293B)
COLOR_CODE = RGBColor(3, 105, 161)        # Cyan/Blue code (#0369A1)

def set_cell_border(cell, **kwargs):
    """Sets borders for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key, val in edge_data.items():
                element.set(qn('w:{}'.format(key)), str(val))

def set_cell_shading(cell, color_hex):
    """Sets cell background color."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def add_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.space_before = Pt(0)
    return p

def add_heading_1(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    return p

def add_heading_2(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = COLOR_GOLD
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    return p

def add_heading_3(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    return p

def add_p(text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Arial"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_PRIMARY
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.font.color.rgb = COLOR_TEXT
    return p

def add_bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Arial"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_PRIMARY
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.font.color.rgb = COLOR_TEXT
    return p

def add_code_block(code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_shading(cell, "F1F5F9")
    set_cell_border(cell,
        top={'val': 'single', 'sz': 4, 'color': 'CBD5E1'},
        bottom={'val': 'single', 'sz': 4, 'color': 'CBD5E1'},
        left={'val': 'single', 'sz': 16, 'color': 'D97706'},
        right={'val': 'single', 'sz': 4, 'color': 'CBD5E1'}
    )
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(code_text)
    run.font.name = "Courier New"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(15, 23, 42)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_screenshot_placeholder(label, caption):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_shading(cell, "F8FAFC")
    set_cell_border(cell,
        top={'val': 'dashed', 'sz': 8, 'color': '94A3B8'},
        bottom={'val': 'dashed', 'sz': 8, 'color': '94A3B8'},
        left={'val': 'dashed', 'sz': 8, 'color': '94A3B8'},
        right={'val': 'dashed', 'sz': 8, 'color': '94A3B8'}
    )
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(24)
    
    r_tag = p.add_run(f"📷 {label}\n")
    r_tag.font.name = "Arial"
    r_tag.font.size = Pt(11)
    r_tag.font.bold = True
    r_tag.font.color.rgb = COLOR_GOLD

    r_hint = p.add_run("[ Reserved screenshot insertion area — 16:9 / 4:3 image ]")
    r_hint.font.name = "Arial"
    r_hint.font.size = Pt(9)
    r_hint.font.italic = True
    r_hint.font.color.rgb = COLOR_MUTED

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(10)
    r_cap_label = p_cap.add_run("Figure: ")
    r_cap_label.font.name = "Arial"
    r_cap_label.font.size = Pt(9)
    r_cap_label.font.bold = True
    r_cap_label.font.color.rgb = COLOR_PRIMARY

    r_cap = p_cap.add_run(caption)
    r_cap.font.name = "Arial"
    r_cap.font.size = Pt(9)
    r_cap.font.italic = True
    r_cap.font.color.rgb = COLOR_TEXT

# ==============================================================================
# REPORT CONTENT GENERATION
# ==============================================================================

# Title & Metadata
add_title("Assignment #6: jQuery\n(Selectors, Effects, Animations)")

p_sub = doc.add_paragraph()
r_sub = p_sub.add_run("Comprehensive Technical Implementation, Verification & Academic Report")
r_sub.font.name = "Arial"
r_sub.font.size = Pt(12)
r_sub.font.bold = True
r_sub.font.color.rgb = COLOR_GOLD
p_sub.paragraph_format.space_after = Pt(12)

# Student Metadata Table
meta_tbl = doc.add_table(rows=5, cols=2)
meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
meta_data = [
    ("Student Name", "Azamat Bukharbayev"),
    ("Academic Group", "IT-2510 (Astana IT University)"),
    ("Selected Task Options", "Task 1: Selector Detective | Task 2: The Party in Motion"),
    ("GitHub Repository", "https://github.com/azamatbukharbayev/Assignment6.git"),
    ("Course & Semester", "Web Technologies 1 (Front End) — Fall 2026")
]
for i, (k, v) in enumerate(meta_data):
    c0 = meta_tbl.cell(i, 0)
    c1 = meta_tbl.cell(i, 1)
    c0.width = Inches(2.2)
    c1.width = Inches(4.3)
    set_cell_shading(c0, "F1F5F9")
    set_cell_shading(c1, "FFFFFF")
    set_cell_border(c0, bottom={'val': 'single', 'sz': 4, 'color': 'E2E8F0'}, right={'val': 'single', 'sz': 4, 'color': 'E2E8F0'})
    set_cell_border(c1, bottom={'val': 'single', 'sz': 4, 'color': 'E2E8F0'})
    p0 = c0.paragraphs[0]
    p0.paragraph_format.space_after = Pt(2)
    p0.paragraph_format.space_before = Pt(2)
    r0 = p0.add_run(k)
    r0.font.bold = True
    r0.font.size = Pt(9.5)
    p1 = c1.paragraphs[0]
    p1.paragraph_format.space_after = Pt(2)
    p1.paragraph_format.space_before = Pt(2)
    r1 = p1.add_run(v)
    r1.font.size = Pt(9.5)

doc.add_paragraph().paragraph_format.space_after = Pt(10)

# ------------------------------------------------------------------------------
# SECTION 1: PROJECT OVERVIEW & TECHNICAL SETUP
# ------------------------------------------------------------------------------
add_heading_1("1. Project Overview & Technical Setup")

add_p(
    "Assignment #6 explores the capabilities of jQuery for advanced DOM selection, structural traversal, animated effects, "
    "and interactive interface engineering. The project is situated around a scenario where AITU students gathered at a birthday party "
    "in a multi-story house between 19:30 and 22:30. At 21:30, the birthday cake mysteriously disappeared from the kitchen. "
    "The assignment is divided into two distinct, independently functioning web applications:"
)

add_bullet("Task 1 (Selector Detective): A data-driven website reconstructing the house, floors, and 21 party guests from the authoritative reference illustration. It implements 20 interactive detective case buttons that execute explicit jQuery selector and traversal chains against the rendered DOM to solve queries, animate found cards with a one-time vertical jump, and dynamically report findings.", "• ")
add_bullet("Task 2 (The Party in Motion): An independent animated version extending the party interface with five core effects: (1) Spotlight and Clear via fadeTo(), (2) Expanding Guest Detail Cards growing outward from avatars with descriptive sentences, (3) Cross-house flying guest room transfers with gap animation, (4) Positional chronological sorting using a jQuery FLIP technique, and (5) Strict animation queue management handling rapid interactions.", "• ")

add_heading_2("1.1. Strict Technical Constraints & Compliance")
add_p(
    "To satisfy the core academic requirements of the university specification, strict engineering boundaries were enforced across both implementations:"
)
add_bullet("Exclusive Library Policy: jQuery v3.7.1 is the sole JavaScript library used. No UI frameworks (Bootstrap, Tailwind JS), component libraries (React, Vue), or external animation libraries (GSAP, Anime.js) were introduced.", "1. ")
add_bullet("No Inline Event Handlers: No inline 'onclick' attributes exist in HTML; all interactions are managed cleanly via jQuery event listeners (.on('click'), .on('keydown')).", "2. ")
add_bullet("Absolute Ban on CSS Transitions and Animations: Neither task1.css nor task2.css contains any 'transition' declarations, '@keyframes' blocks, or CSS animation properties. All dynamic movement, jumping, opacity changes, and resizing are governed exclusively by jQuery effects (animate(), fadeIn(), fadeOut(), fadeTo()).", "3. ")
add_bullet("DOM-Based Querying Integrity: Selector detective cases evaluate against the live rendered HTML document tree using DOM classes and data-* attributes. The JavaScript array is strictly used for initial rendering and state synchronization, never as a shortcut substitute for DOM querying.", "4. ")

add_heading_2("1.2. Project Architecture & File Organization")
add_p(
    "The project maintains clean modularity and separation of concerns. The production workspace is organized as follows:"
)

add_code_block(
"""Assignment6/
├── index.html                   # Project landing hub & documentation overview
├── task1.html                   # Task 1: Selector Detective web application
├── task2.html                   # Task 2: The Party in Motion web application
├── css/
│   ├── task1.css                # Task 1 static styles (zero transitions/animations)
│   └── task2.css                # Task 2 static styles & modal overlays
├── js/
│   ├── jquery.min.js            # Standalone local jQuery 3.7.1 dependency
│   ├── task1.js                 # Task 1 data model, DOM renderer, 20 selector cases
│   └── task2.js                 # Task 2 independent mutable model, effects, sorting
├── test_runner.js               # Headless verification suite (Node.js + jsdom)
├── generate_report.py           # Automated academic report generator
├── README.md                    # Operational manual, test guide, defense notes
├── Assignment6_AzamatBukharbayev.docx # Editable academic report
└── Assignment6_AzamatBukharbayev.pdf  # Final rendered PDF report"""
)

# ------------------------------------------------------------------------------
# SECTION 2: RECONSTRUCTING THE ORIGINAL PARTY DATA
# ------------------------------------------------------------------------------
add_heading_1("2. Reconstruction of the Original Party Data")

add_p(
    "The starting point of the assignment is the precise transcription of the reference illustration ('The Party, 19:30 - 22:30'). "
    "Every visible detail in the image—including the house's physical architecture, the vertical floor distribution, the horizontal "
    "left-to-right order of guests, guest names, clothing ring colors, arrival timestamps, AITU student affiliations, and every "
    "associated item—was transcribed into a structured JavaScript model without omission or modification."
)

add_heading_2("2.1. House Geometry & Floor Distribution")
add_p(
    "The house consists of two distinct floors and four rooms:"
)
add_bullet("Floor 2 (Upper Floor): Contains Balcony (left, 4 guests) and Hall (right, 6 guests).", "• ")
add_bullet("Floor 1 (Ground Floor): Contains Kitchen (left, 6 guests) and Living Room (right, 5 guests).", "• ")
add_p("The initial party comprises exactly 21 guests across 4 rooms.")

add_heading_2("2.2. Authoritative Initial Party Dataset")
add_p(
    "Table 1 outlines the complete reconstructed party dataset used as the immutable baseline in Task 1 and initial state in Task 2:"
)

# Party Data Table
data_tbl = doc.add_table(rows=22, cols=8)
data_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ["Room", "Floor", "Name", "Clothing Ring", "Arrival", "AITU?", "Items Associated", "Init. Order"]
for j, h in enumerate(headers):
    cell = data_tbl.cell(0, j)
    set_cell_shading(cell, "0F172A")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(h)
    r.font.bold = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(255, 255, 255)

guests_raw = [
    # Balcony
    ("Balcony", "2", "Aigerim", "Green", "20:45", "Yes", "Cup, Cake crumbs", "0"),
    ("Balcony", "2", "Aliya", "Black", "21:15", "No", "Phone, Cup", "1"),
    ("Balcony", "2", "Madina", "Green", "22:10", "No", "Phone, Bag, Cake crumbs", "2"),
    ("Balcony", "2", "Ivan", "Blue", "20:10", "Yes", "Cup", "3"),
    # Hall
    ("Hall", "2", "Arman", "Yellow", "20:35", "No", "Bag, Glasses", "0"),
    ("Hall", "2", "Sophia", "Purple", "21:20", "Yes", "Phone", "1"),
    ("Hall", "2", "Dias", "Green", "19:45", "Yes", "None (empty)", "2"),
    ("Hall", "2", "Bekzat", "Purple", "20:55", "No", "Bag, Cup, Cake crumbs", "3"),
    ("Hall", "2", "Aruzhan", "Blue", "21:50", "Yes", "Phone", "4"),
    ("Hall", "2", "Saule", "Green", "21:25", "No", "Glasses", "5"),
    # Kitchen
    ("Kitchen", "1", "Nurlan", "Red", "19:40", "Yes", "Phone, Cup", "0"),
    ("Kitchen", "1", "Emma", "Black", "20:15", "Yes", "Bag", "1"),
    ("Kitchen", "1", "Timur", "Green", "21:05", "Yes", "Phone, Bag, Cake crumbs", "2"),
    ("Kitchen", "1", "Sergey", "Yellow", "20:50", "No", "None (empty)", "3"),
    ("Kitchen", "1", "Yerlan", "Red", "21:40", "Yes", "Glasses, Phone", "4"),
    ("Kitchen", "1", "Kamila", "Green", "20:05", "No", "Cup", "5"),
    # Living Room
    ("Living room", "1", "Daniel", "Blue", "19:55", "Yes", "Glasses", "0"),
    ("Living room", "1", "Zarina", "Green", "21:10", "No", "Bag, Cup, Cake crumbs", "1"),
    ("Living room", "1", "Natasha", "Yellow", "20:30", "Yes", "None (empty)", "2"),
    ("Living room", "1", "Askar", "Blue", "21:00", "Yes", "Phone", "3"),
    ("Living room", "1", "Inkar", "Purple", "20:20", "No", "Bag", "4"),
]

for i, row_data in enumerate(guests_raw):
    for j, val in enumerate(row_data):
        cell = data_tbl.cell(i + 1, j)
        if i % 2 == 1:
            set_cell_shading(cell, "F8FAFC")
        set_cell_border(cell, bottom={'val': 'single', 'sz': 4, 'color': 'E2E8F0'})
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(val)
        r.font.size = Pt(8.5)
        if j in (0, 1, 4, 5, 7):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if j == 2:
            r.font.bold = True

doc.add_paragraph().paragraph_format.space_after = Pt(8)

add_heading_2("2.3. DOM Construction & Avatar Integration")
add_p(
    "The party house is assembled dynamically using jQuery DOM construction (createGuestCard). "
    "For each guest, an interactive card element (.guest-card) is created containing:"
)
add_bullet("Avatar Generation: SVGs sourced from DiceBear Open Peeps (https://api.dicebear.com/9.x/open-peeps/svg?seed=GuestName) encoded safely via encodeURIComponent. If an image fails to load, a resilient inline fallback avatar is substituted.", "1. ")
add_bullet("Clothing Outline Ring: A 3.5px solid circular border wrapping the avatar with specific color classes (.ring-red, .ring-blue, .ring-green, .ring-yellow, .ring-black, .ring-purple).", "2. ")
add_bullet("AITU Student Badge: A graduation cap badge (🎓) positioned at the top-right of the avatar ring for students having aituStudent: true.", "3. ")
add_bullet("Arrival Timestamp: Displayed with tabular numbers alongside a clock icon (⏱️ HH:MM).", "4. ")
add_bullet("Item Representation: Specific CSS classes (.item-phone, .item-glasses, .item-bag, .item-cup, .item-crumbs) and lightweight accessible text icons with tooltip descriptions. Guests with no items receive class .has-no-items.", "5. ")
add_bullet("Data Attributes: Attributes data-id, data-name, data-color, data-arrival, data-aitu, data-room, and data-floor are populated, enabling compound DOM queries.", "6. ")

# ------------------------------------------------------------------------------
# SECTION 3: TASK 1 — SELECTOR DETECTIVE (20 CASES)
# ------------------------------------------------------------------------------
add_heading_1("3. Task 1 — Selector Detective (All 20 Cases)")

add_p(
    "In Task 1, all 20 detective cases were implemented using jQuery selector chains and traversal methods executed against the rendered HTML DOM. "
    "Before running any case (1–17, 20), the house is freshly reconstructed from the original data array to ensure complete test isolation. "
    "Cases 18 and 19 represent dynamic mutations, adding Miras via fadeIn() and removing Natasha via fadeOut()."
)

cases_documentation = [
    (
        1,
        "Who is in the hall?",
        '$("#room-hall .guest-card")',
        "Uses an ID descendant selector to select all guest cards located within the room element whose id is 'room-hall'. Matches all 6 residents currently standing in the Hall.",
        "6 guests",
        "Arman, Sophia, Dias, Bekzat, Aruzhan, Saule",
        "TASK 1 / CASE 1",
        "Running party page immediately after Case 1 execution. The 6 Hall guests are highlighted with yellow glowing borders and have jumped once with animate(). The result banner displays count 6 and their respective names."
    ),
    (
        2,
        "Who is dressed in yellow or purple?",
        '$(".guest-card.yellow, .guest-card.purple")',
        "Employs a multiple selector combinator (comma-separated) targeting guest cards possessing either the .yellow or .purple clothing class. Retrieves 3 yellow-clothed and 3 purple-clothed guests.",
        "6 guests",
        "Arman, Sophia, Bekzat, Sergey, Natasha, Inkar",
        "TASK 1 / CASE 2",
        "Guest cards wearing yellow or purple rings highlighted across all rooms. Result area shows count 6 and guest names."
    ),
    (
        3,
        "Which AITU students are upstairs?",
        '$(".floor-section[data-floor=\'2\'] .guest-card.aitu")',
        "Combines an attribute selector targeting the second-floor container (.floor-section[data-floor='2']) with a class selector for AITU students (.guest-card.aitu). Filters out first-floor students and non-AITU upstairs guests.",
        "5 guests",
        "Aigerim, Ivan, Sophia, Dias, Aruzhan",
        "TASK 1 / CASE 3",
        "Upstairs AITU students highlighted on Floor 2. Excludes downstairs students (Nurlan, Timur, Yerlan, Daniel, Natasha, Askar)."
    ),
    (
        4,
        "Who has glasses or a bag?",
        '$(".guest-card:has(.item-glasses), .guest-card:has(.item-bag)")',
        "Leverages the jQuery :has() pseudo-selector within a multiple selector to identify cards containing either glasses (.item-glasses) or a bag (.item-bag). Guests possessing both items (e.g., Arman) appear exactly once.",
        "10 guests",
        "Madina, Arman, Bekzat, Saule, Emma, Timur, Yerlan, Daniel, Zarina, Inkar",
        "TASK 1 / CASE 4",
        "Guests possessing glasses or bags highlighted. Verifies that cards with both items are not duplicated."
    ),
    (
        5,
        "Who came with nothing at all?",
        '$(".guest-card:not(:has(.item-icon))")',
        "Combines the negation pseudo-class :not() with the hierarchical filter :has(.item-icon). Accurately distinguishes guests with zero associated items from those carrying accessories, ignoring avatar and badges.",
        "3 guests",
        "Dias, Sergey, Natasha",
        "TASK 1 / CASE 5",
        "The three guests who arrived with no accessories (Dias, Sergey, Natasha) highlighted with count 3."
    ),
    (
        6,
        "Who has a phone but no bag?",
        '$(".guest-card:has(.item-phone):not(:has(.item-bag))")',
        "Chains :has(.item-phone) with :not(:has(.item-bag)) to enforce a strict logical conjunction: presence of a phone and total absence of a bag. Excludes guests carrying both (Madina, Timur).",
        "6 guests",
        "Aliya, Sophia, Aruzhan, Nurlan, Yerlan, Askar",
        "TASK 1 / CASE 6",
        "Guests with a phone and no bag highlighted across rooms. Excludes Madina and Timur."
    ),
    (
        7,
        "Who arrived between 20:30 and 21:00, both times included?",
        '$(".guest-card").filter(function() { const t = $(this).attr("data-arrival"); return t >= "20:30" && t <= "21:00"; })',
        "Applies jQuery .filter() passing a predicate function comparing the data-arrival attribute. Inclusive inequality (>= '20:30' && <= '21:00') guarantees boundary captures (Natasha at 20:30, Askar at 21:00).",
        "6 guests",
        "Aigerim (20:45), Arman (20:35), Bekzat (20:55), Sergey (20:50), Natasha (20:30), Askar (21:00)",
        "TASK 1 / CASE 7",
        "Guests arriving in the inclusive 20:30–21:00 window highlighted. Natasha (20:30) and Askar (21:00) verified."
    ),
    (
        8,
        "Who arrived after 21:00 and is not an AITU student?",
        '$(".guest-card:not(.aitu)").filter(function() { return $(this).attr("data-arrival") > "21:00"; })',
        "Pre-filters non-AITU guests using :not(.aitu), then evaluates strict inequality (> '21:00') over data-arrival. Excludes guests arriving at exactly 21:00 (Askar) and late-arriving AITU students.",
        "4 guests",
        "Aliya (21:15), Madina (22:10), Saule (21:25), Zarina (21:10)",
        "TASK 1 / CASE 8",
        "Non-AITU guests who arrived strictly after 21:00 highlighted. Count 4 visible in result area."
    ),
    (
        9,
        "Who stands right next to Askar?",
        '$(".guest-card[data-name=\'Askar\']").prev(".guest-card").add($(".guest-card[data-name=\'Askar\']").next(".guest-card"))',
        "Locates Askar via an attribute selector, then combines his immediate predecessor (.prev()) and immediate successor (.next()) using .add(). Excludes Askar himself and non-immediate room occupants.",
        "2 guests",
        "Natasha (left), Inkar (right)",
        "TASK 1 / CASE 9",
        "Askar's immediate physical neighbours in the Living Room highlighted without selecting Askar."
    ),
    (
        10,
        "Who stands last in each room?",
        '$(".room .guest-card:last-child")',
        "Utilizes the CSS structural pseudo-class :last-child scoped to room containers. Finds the rightmost card rendered in each of the four rooms.",
        "4 guests",
        "Ivan (Balcony), Saule (Hall), Kamila (Kitchen), Inkar (Living room)",
        "TASK 1 / CASE 10",
        "The final card in each of the four rooms highlighted. Count 4 displayed."
    ),
    (
        11,
        "Who stands third in each room?",
        '$(".room .guest-card:nth-child(3)")',
        "Applies :nth-child(3) scoped to each room's guest container, selecting the card at index 3 (1-based index) in left-to-right order.",
        "4 guests",
        "Madina (Balcony), Dias (Hall), Timur (Kitchen), Natasha (Living room)",
        "TASK 1 / CASE 11",
        "The third card in each room highlighted across Balcony, Hall, Kitchen, and Living Room."
    ),
    (
        12,
        "Who stands after Timur in his room?",
        '$(".guest-card[data-name=\'Timur\']").nextAll(".guest-card")',
        "Demonstrates the .nextAll() traversal method starting from Timur's card, capturing all subsequent siblings in the Kitchen. Returns Sergey, Yerlan, and Kamila.",
        "3 guests",
        "Sergey, Yerlan, Kamila",
        "TASK 1 / CASE 12",
        "All guests positioned after Timur in the Kitchen highlighted via nextAll(). Count 3 visible."
    ),
    (
        13,
        "Who stands right before someone in black?",
        '$(".guest-card.black").prev(".guest-card")',
        "Selects all guests wearing black (.black), then invokes .prev('.guest-card') to retrieve their immediate left predecessors within the same room container.",
        "2 guests",
        "Aigerim (before Aliya in Balcony), Nurlan (before Emma in Kitchen)",
        "TASK 1 / CASE 13",
        "Immediate predecessors of black-clothed guests highlighted. Aigerim and Nurlan identified."
    ),
    (
        14,
        "In which room is Aruzhan?",
        '$(".guest-card[data-name=\'Aruzhan\']").closest(".room")',
        "Traverses upward from Aruzhan's card to its parent room container using .closest('.room'). Highlights the Hall room element and reports room identity rather than guest card.",
        "1 room",
        "Hall (Floor 2)",
        "TASK 1 / CASE 14",
        "The Hall room container highlighted with yellow border; guests inside Hall jump once. Result banner displays '1 room: Hall'."
    ),
    (
        15,
        "Which rooms have fewer than 6 guests?",
        '$(".room").filter(function() { return $(this).find(".guest-card").length < 6; })',
        "Filters the four room elements using .filter() by evaluating the length of $(this).find('.guest-card'). Balcony has 4 guests and Living Room has 5 guests, satisfying < 6.",
        "2 rooms",
        "Balcony (4 guests), Living room (5 guests)",
        "TASK 1 / CASE 15",
        "Balcony and Living room highlighted. Count 2 rooms displayed."
    ),
    (
        16,
        "In which rooms does somebody have a cup, but nobody wears glasses?",
        '$(".room:has(.item-cup):not(:has(.item-glasses))")',
        "Evaluates rooms possessing at least one cup (:has(.item-cup)) while strictly having zero pairs of glasses (:not(:has(.item-glasses))). Evaluates independently per room; only Balcony satisfies both.",
        "1 room",
        "Balcony (3 cups, 0 glasses)",
        "TASK 1 / CASE 16",
        "Balcony room highlighted as the sole room meeting the compound constraint. Result area shows Balcony."
    ),
    (
        17,
        "Who came first to each room?",
        '$(".room").map(function() { let e = null, min = "99:99"; $(this).find(".guest-card").each(function() { const t = $(this).attr("data-arrival"); if (t < min) { min = t; e = this; } }); return e; })',
        "Iterates over rooms, evaluating data-arrival across child cards to find the chronological minimum timestamp in each room. Returns Ivan (20:10), Dias (19:45), Nurlan (19:40), and Daniel (19:55).",
        "4 guests",
        "Ivan (20:10), Dias (19:45), Nurlan (19:40), Daniel (19:55)",
        "TASK 1 / CASE 17",
        "The earliest arrival in every room highlighted. Real timestamps parsed and compared."
    ),
    (
        18,
        "Miras joins Askar and Inkar",
        'const mirasCard = createGuestCard(mirasData).hide(); mirasCard.insertAfter($(".guest-card[data-name=\'Askar\']")).fadeIn(600);',
        "Rebuilds the initial party, instantiates a valid Miras card (21:20, blue clothing, AITU badge, phone), inserts him between Askar and Inkar using insertAfter(), and smoothly reveals him with fadeIn(600).",
        "1 guest added",
        "Miras (Living room, between Askar and Inkar)",
        "TASK 1 / CASE 18",
        "Miras card rendered between Askar and Inkar in the Living Room following smooth fadeIn(600). Card exhibits blue ring, 21:20 arrival, phone icon, and AITU badge."
    ),
    (
        19,
        "Natasha leaves the party",
        'const natasha = $(".guest-card[data-name=\'Natasha\']"); natasha.fadeOut(600, function() { $(this).remove(); });',
        "Rebuilds initial party, locates Natasha's card, applies a goodbye highlight, animates her departure via fadeOut(600), and removes the DOM element upon completion using .remove().",
        "1 guest departed",
        "Natasha (card removed from Living room)",
        "TASK 1 / CASE 19",
        "Living Room after Natasha's departure. Card completely removed from DOM; no ghost element remains."
    ),
    (
        20,
        "Who took the cake?",
        '$(".floor-section[data-floor=\'1\'] .guest-card").filter(":has(.item-crumbs):not(:has(.item-cup))").filter(function() { const t = $(this).attr("data-arrival"); return t > "21:00" && t < "21:30"; }).filter(function() { return $(this).prev(".aitu").length > 0 || $(this).next(".aitu").length > 0; })',
        "Evaluates the complete 4-clue logical conjunction against the DOM: (1) Downstairs (.floor-section[data-floor='1']), (2) Strict arrival window 21:00 < t < 21:30, (3) Possesses cake crumbs without a cup, and (4) Stood immediately adjacent to an AITU student (.prev('.aitu') or .next('.aitu')). Uniquely unmasks Timur.",
        "1 thief identified",
        "Timur (Kitchen, Floor 1)",
        "TASK 1 / CASE 20",
        "Timur highlighted in the Kitchen as the cake thief. Result banner explains how all four clues converge strictly and uniquely on him."
    )
]

for case_item in cases_documentation:
    c_num, c_q, c_chain, c_exp, c_cnt, c_res, p_lbl, p_cap = case_item
    add_heading_2(f"Case {c_num} — {c_q}")
    add_p(c_exp)
    add_p("Actual jQuery Selector Chain:", bold_prefix="Code Snippet: ")
    add_code_block(c_chain)
    add_p(f"Count: {c_cnt} | Identified: {c_res}", bold_prefix="Verified Result: ")
    add_screenshot_placeholder(f"[SCREENSHOT PLACEHOLDER — {p_lbl}]", p_cap)

# ------------------------------------------------------------------------------
# SECTION 4: TASK 2 — THE PARTY IN MOTION
# ------------------------------------------------------------------------------
add_heading_1("4. Task 2 — The Party in Motion (Effects & Animations)")

add_p(
    "Task 2 develops an independent copy of the party web application that brings the interface to life through interactive jQuery animations. "
    "Unlike Task 1, Task 2 preserves mutable party state: moving a guest updates both the live DOM and the underlying data array, "
    "so that subsequent queries reflect the moved guest's new location. CSS transitions and CSS animations remain strictly prohibited."
)

add_heading_2("4.1. Effect 1 — Spotlight & Clear (fadeTo)")
add_p(
    "Clicking any case button highlights matching elements while dimming unrelated cards. Target cards remain fully opaque (opacity: 1), "
    "while non-target cards smoothly fade to 20% opacity over 300 milliseconds using jQuery's fadeTo(300, 0.2). "
    "When a room query executes (Cases 14–16), all guests in the matched rooms remain bright, and guests in non-matching rooms fade. "
    "Clicking 'Clear Spotlight' executes fadeTo(300, 1.0) across all guest cards and removes highlight borders."
)

add_screenshot_placeholder(
    "[SCREENSHOT PLACEHOLDER — TASK 2 / SPOTLIGHT]",
    "Running Task 2 after clicking Case 3. Upstairs AITU students on Floor 2 remain at full opacity (1.0), while all other guest cards across the house fade to 20% opacity (0.2)."
)

add_screenshot_placeholder(
    "[SCREENSHOT PLACEHOLDER — TASK 2 / CLEAR]",
    "Running Task 2 after clicking Clear Spotlight. All guest cards are restored to full opacity (1.0) and highlight styling is cleared."
)

add_heading_2("4.2. Effect 2 — Expanding Guest Detail Card")
add_p(
    "Clicking any guest card dynamically spawns an interactive modal card displaying full guest information formatted in descriptive words: "
    "arrival time ('Came at 20:30.'), student status ('AITU student.'), accessories ('Has a phone and a bag.'), clothing ('Wears yellow.'), "
    "and current room ('Currently in the Living room (Floor 1).')."
)
add_bullet("Opening Animation: The card does not abruptly appear. Using $avatar.offset(), the modal begins at the avatar's exact coordinates and dimensions (48x48px, opacity 0.2) and expands outward to centered modal dimensions (360px wide, opacity 1.0) over 300ms via jQuery animate().", "• ")
add_bullet("Closing Animation: Triggered when pressing Escape, clicking the backdrop outside the card, or clicking the close button. The card animates shrinking back into the originating avatar's coordinates over 240ms, safely removing the element upon completion.", "• ")

add_screenshot_placeholder(
    "[SCREENSHOT PLACEHOLDER — TASK 2 / EXPANDED GUEST]",
    "Expanded guest detail card opened after clicking Arman. Demonstrates outward expansion from his avatar, clear descriptive sentences, and room movement buttons."
)

add_screenshot_placeholder(
    "[SCREENSHOT PLACEHOLDER — TASK 2 / ESCAPE CLOSE]",
    "View immediately after pressing Escape. The detail card has shrunk back into Arman's avatar and the overlay is cleanly removed from the DOM."
)

add_heading_2("4.3. Effect 3 — Cross-House Flying Room Transfers")
add_p(
    "Inside the guest detail card, action buttons allow moving the guest to any other room (e.g., 'Move to Balcony'). "
    "When clicked, the transfer executes as a continuous physical flight rather than an abrupt teleportation:"
)
add_bullet("Placeholder Reservation: An animated placeholder (.guest-card-placeholder) is appended to the target room's list, expanding its width from 0 to 78px over 280ms to smoothly open space.", "1. ")
add_bullet("Flight Animation: A visual clone (.flying-guest-card) is created at the card's starting offset and animated across the screen via animate({ top: dest.top, left: dest.left }, 500). Simultaneously, the original card in the old room animates its width to 0, smoothly closing the vacated space.", "2. ")
add_bullet("State Synchronization: Upon arrival, the clone and placeholder are removed, the real card is placed at the end of the destination room with updated data-room and data-floor attributes, and the guest object is transferred between arrays in the liveParty data structure.", "3. ")

add_screenshot_placeholder(
    "[SCREENSHOT PLACEHOLDER — TASK 2 / GUEST MOVEMENT]",
    "Arman positioned at the right end of the Balcony following cross-house flight. The Hall gap is completely closed, Balcony room badge reads '5 guests', and Arman's attributes reflect Floor 2 Balcony."
)

add_screenshot_placeholder(
    "[SCREENSHOT PLACEHOLDER — TASK 2 / CASE 1 AFTER MOVEMENT]",
    "Execution of Case 1 (Who is in the hall?) after Arman's transfer to Balcony. Hall result highlights 5 guests and no longer contains Arman."
)

add_heading_2("4.4. Effect 4 — Chronological Positional Sorting")
add_p(
    "Each room header contains a 'Sort by time' button. Clicking it reorders that room's cards chronologically by arrival time (earliest to latest). "
    "Clicking it a second time restores the original picture order. The animation uses a jQuery positional FLIP technique:"
)
add_bullet("First: Capture the starting offset() of every guest card in the room.", "1. ")
add_bullet("Last: Reorder the cards in the DOM using jQuery append() and capture the new offset().", "2. ")
add_bullet("Invert: Apply relative positional deltas (left = oldLeft - newLeft, top = oldTop - newTop) via .css() so cards visually remain at their old locations.", "3. ")
add_bullet("Play: Animate left and top to 0 over 380ms via animate(), producing smooth sliding movement across positions without jumping.", "4. ")

add_screenshot_placeholder(
    "[SCREENSHOT PLACEHOLDER — TASK 2 / SORTED LIVING ROOM]",
    "Living Room after clicking 'Sort by time'. Guests slide smoothly into chronological order: Daniel (19:55), Inkar (20:20), Natasha (20:30), Askar (21:00), Zarina (21:10)."
)

add_screenshot_placeholder(
    "[SCREENSHOT PLACEHOLDER — TASK 2 / RESTORED ORDER]",
    "Living Room after clicking 'Sort by time' a second time. Cards slide smoothly back to their original picture order: Daniel, Zarina, Natasha, Askar, Inkar."
)

add_heading_2("4.5. Effect 5 — Animation Queue & Fast-Click Parity")
add_p(
    "To prevent animation queue buildup, element duplication, and race conditions during rapid user interaction, "
    "strict queue management was implemented using jQuery's .stop(true, false), operation tokens, and parity tracking. "
    "When a user clicks the sort button five times in rapid succession, the state toggles an odd number of times (5), "
    "guaranteeing that the room cleanly settles in sorted order without shaking, jumping, or leaving residual animation queues."
)

add_screenshot_placeholder(
    "[SCREENSHOT PLACEHOLDER — TASK 2 / RAPID SORTING]",
    "Living Room after five rapid sorting clicks. The room settles cleanly in chronological sorted order with zero duplicate cards, stuck queues, or displaced elements."
)

# Acceptance Test Table
add_heading_2("4.6. Authoritative Task 2 Verification Table")
add_p(
    "Table 2 documents the five authoritative acceptance tests specified in the assignment, comparing expected results against actual verified behavior:"
)

test_tbl = doc.add_table(rows=6, cols=5)
test_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
test_headers = ["Test Check", "Procedure", "Expected Result", "Actual Verified Result", "Status"]
for j, h in enumerate(test_headers):
    cell = test_tbl.cell(0, j)
    set_cell_shading(cell, "0F172A")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(h)
    r.font.bold = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(255, 255, 255)

acceptance_tests = [
    (
        "Case 3, then Clear",
        "Click Case 3 button, wait 350ms, then click Clear Spotlight.",
        "Only AITU students on 2nd floor stay bright; Clear brings everybody back to full opacity.",
        "Verified: Exactly 5 upstairs AITU students remained at opacity 1.0 while others faded to 0.2 over 300ms. Clear restored all 21 cards to opacity 1.0.",
        "Pass"
    ),
    (
        "Click a guest, then Esc",
        "Click any guest card (Aigerim), observe expansion, then press Escape.",
        "The card grows out of the avatar with all data; Esc shrinks it back into the avatar.",
        "Verified: Modal expanded smoothly from avatar offset (300ms) with full descriptive sentences; Escape key triggered 240ms shrink animation back into avatar.",
        "Pass"
    ),
    (
        "Move Arman to balcony",
        "Open Arman's card, click 'Move to Balcony', wait for animation, then execute Case 1.",
        "Arman flies to end of balcony; hall closes the gap; Case 1 no longer finds him.",
        "Verified: Arman animated across house (500ms), Hall gap closed (280ms), Balcony opened space. DOM & array updated; subsequent Case 1 query excludes Arman.",
        "Pass"
    ),
    (
        "Sort living room twice",
        "Click 'Sort by time' in Living room header, observe order, then click again.",
        "Cards slide to 19:55, 20:20, 20:30, 21:00, 21:10, then slide back to picture order.",
        "Verified: Cards smoothly slid into chronological order (19:55 → 20:20 → 20:30 → 21:00 → 21:10), second click restored Daniel → Zarina → Natasha → Askar → Inkar.",
        "Pass"
    ),
    (
        "Sort 5 times fast",
        "Click 'Sort by time' in Living room 5 times in rapid succession.",
        "Cards end in sorted order without jumping or shaking; no animation queues repeat.",
        "Verified: Handled odd-parity toggles cleanly. Finished in sorted order (19:55 to 21:10) without shaking, duplicate cards, or continuous animation queues.",
        "Pass"
    )
]

for i, test_data in enumerate(acceptance_tests):
    for j, val in enumerate(test_data):
        cell = test_tbl.cell(i + 1, j)
        if i % 2 == 1:
            set_cell_shading(cell, "F8FAFC")
        set_cell_border(cell, bottom={'val': 'single', 'sz': 4, 'color': 'E2E8F0'})
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(val)
        r.font.size = Pt(8.5)
        if j == 4:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r.font.bold = True
            r.font.color.rgb = RGBColor(22, 163, 74) # Green
        elif j == 0:
            r.font.bold = True

doc.add_paragraph().paragraph_format.space_after = Pt(10)

# ------------------------------------------------------------------------------
# SECTION 5: CONCLUSION & TECHNICAL INSIGHTS
# ------------------------------------------------------------------------------
add_heading_1("5. Conclusion & Technical Insights")

add_p(
    "Completing Assignment #6 has significantly advanced my practical mastery of core frontend engineering, DOM manipulation, "
    "and animation sequencing using jQuery. By strictly avoiding external frameworks, CSS transitions, and CSS keyframe animations, "
    "I was required to solve complex motion challenges purely through algorithmic coordinate calculations and jQuery's effect APIs."
)

add_p(
    "In Task 1, constructing the party house from a data model reinforced how data-* attributes and semantic class names "
    "bridge JavaScript data structures with DOM query engines. Implementing the 20 selector detective cases taught me the nuances "
    "between hierarchical descendant selectors, jQuery-specific extensions (:has(), :not()), structural position selectors (:last-child, :nth-child()), "
    "and traversal methods (.next(), .prev(), .nextAll(), .closest(), .add()). Solving Case 20 (the cake mystery) through a compound four-clue filter "
    "demonstrated the power of functional DOM filtering when analyzing complex relational conditions."
)

add_p(
    "In Task 2, building the animated party interface revealed the critical importance of keeping application state synchronized across both "
    "the live DOM and memory arrays. Engineering the cross-house flying transfer required coordinating offset calculations, temporary clones, "
    "and simultaneous dimension animations so that moving elements feel physical rather than synthetic. Similarly, implementing chronological sorting "
    "via the FLIP technique demonstrated how measuring starting and ending geometry allows smooth positional sliding without layout jitter. "
    "Finally, handling rapid user clicks with .stop(true, false) and parity tracking provided essential lessons in concurrency control "
    "and animation queue hygiene, ensuring the interface remains deterministic, responsive, and robust under all conditions."
)

doc.save(docx_path)
print(f"DOCX report saved successfully at: {docx_path}")

# ==============================================================================
# PDF EXPORT ATTEMPT (Apple Pages or ReportLab)
# ==============================================================================
print("Attempting PDF export...")

# Strategy 1: Apple Pages via AppleScript
applescript_code = f'''
tell application "Pages"
    set myDoc to open (POSIX file "{docx_path}")
    export myDoc to (POSIX file "{pdf_path}") as PDF
    close myDoc saving no
end tell
'''

try:
    res = subprocess.run(["osascript", "-e", applescript_code], capture_output=True, text=True, timeout=30)
    if res.returncode == 0 and os.path.exists(pdf_path) and os.path.getsize(pdf_path) > 1000:
        print(f"PDF exported successfully via Apple Pages at: {pdf_path}")
        sys.exit(0)
    else:
        print(f"Pages AppleScript export returned code {res.returncode}: {res.stderr}")
except Exception as e:
    print(f"Pages AppleScript export failed: {e}")

# Strategy 2: ReportLab fallback if Pages was blocked
print("Running ReportLab PDF generator fallback...")
try:
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Preformatted
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors

    pdf_doc = SimpleDocTemplate(pdf_path, pagesize=letter, rightMargin=54, leftMargin=54, topMargin=54, bottomMargin=54)
    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle('TitleStyle', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=20, leading=24, textColor=colors.HexColor('#0F172A'))
    sub_style = ParagraphStyle('SubStyle', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=11, leading=14, textColor=colors.HexColor('#D97706'))
    h1_style = ParagraphStyle('H1Style', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=14, leading=18, textColor=colors.HexColor('#0F172A'), spaceBefore=14, spaceAfter=6, keepWithNext=True)
    h2_style = ParagraphStyle('H2Style', parent=styles['Heading2'], fontName='Helvetica-Bold', fontSize=11, leading=15, textColor=colors.HexColor('#D97706'), spaceBefore=10, spaceAfter=4, keepWithNext=True)
    p_style = ParagraphStyle('PStyle', parent=styles['Normal'], fontName='Helvetica', fontSize=9.5, leading=13.5, textColor=colors.HexColor('#1E293B'), spaceAfter=5)
    bullet_style = ParagraphStyle('BStyle', parent=p_style, leftIndent=12, bulletIndent=4, spaceAfter=3)
    code_style = ParagraphStyle('CodeStyle', parent=styles['Code'], fontName='Courier', fontSize=8, leading=10, textColor=colors.HexColor('#0F172A'))

    story = []

    story.append(Paragraph("Assignment #6: jQuery (Selectors, Effects, Animations)", title_style))
    story.append(Paragraph("Comprehensive Technical Implementation & Verification Report", sub_style))
    story.append(Spacer(1, 10))

    # Meta table
    meta_table_data = [
        [Paragraph("<b>Student Name</b>", p_style), Paragraph("Azamat Bukharbayev", p_style)],
        [Paragraph("<b>Academic Group</b>", p_style), Paragraph("IT-2510 (Astana IT University)", p_style)],
        [Paragraph("<b>Selected Tasks</b>", p_style), Paragraph("Task 1: Selector Detective | Task 2: The Party in Motion", p_style)],
        [Paragraph("<b>GitHub Repository</b>", p_style), Paragraph("https://github.com/azamatbukharbayev/Assignment6.git", p_style)],
        [Paragraph("<b>Course</b>", p_style), Paragraph("Web Technologies 1 (Front End) — Fall 2026", p_style)]
    ]
    meta_table = Table(meta_table_data, colWidths=[140, 360])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#F1F5F9')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 12))

    # Section 1
    story.append(Paragraph("1. Project Overview & Technical Setup", h1_style))
    story.append(Paragraph("Assignment #6 explores the capabilities of jQuery for advanced DOM selection, structural traversal, animated effects, and interactive interface engineering. Built under strict university guidelines enforcing jQuery as the sole library with zero CSS transitions and zero CSS animations.", p_style))
    story.append(Paragraph("<b>Task 1 (Selector Detective):</b> 20 interactive cases querying the house DOM, triggering one-time jump animations, and dynamically reporting results.", bullet_style))
    story.append(Paragraph("<b>Task 2 (The Party in Motion):</b> Independent animated copy featuring Spotlights (fadeTo), expanding guest cards, cross-house flying room transfers, and chronological positional sorting.", bullet_style))

    # Section 2
    story.append(Paragraph("2. Reconstruction of Original Party Data", h1_style))
    story.append(Paragraph("The party house consists of 2 floors, 4 rooms (Balcony, Hall, Kitchen, Living Room), and exactly 21 initial guests transcribed with 100% precision from the reference illustration.", p_style))

    # Section 3
    story.append(Paragraph("3. Task 1 — Selector Detective (All 20 Cases)", h1_style))
    story.append(Paragraph("All 20 cases are executed against the rendered DOM tree using explicit jQuery selector chains:", p_style))

    for c_num, c_q, c_chain, c_exp, c_cnt, c_res, p_lbl, p_cap in cases_documentation:
        story.append(Paragraph(f"<b>Case {c_num}: {c_q}</b>", h2_style))
        story.append(Paragraph(c_exp, p_style))
        story.append(Paragraph(f"<b>jQuery Chain:</b> <code>{c_chain}</code>", p_style))
        story.append(Paragraph(f"<b>Verified Result:</b> {c_cnt} ({c_res})", p_style))
        # Placeholder box
        box_data = [[Paragraph(f"📷 <b>[{p_lbl}]</b><br/><i>Figure: {p_cap}</i>", p_style)]]
        box_tbl = Table(box_data, colWidths=[500])
        box_tbl.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#94A3B8'), None, (3, 3)),
            ('TOPPADDING', (0, 0), (-1, -1), 8),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER')
        ]))
        story.append(box_tbl)
        story.append(Spacer(1, 6))

    # Section 4
    story.append(Paragraph("4. Task 2 — The Party in Motion (Effects & Animations)", h1_style))
    story.append(Paragraph("Task 2 implements the four animated effects and queue management, verified across five authoritative acceptance tests:", p_style))

    # Acceptance test table in PDF
    pdf_test_data = [
        [Paragraph("<b>Check</b>", p_style), Paragraph("<b>Expected Result</b>", p_style), Paragraph("<b>Actual Verified Result</b>", p_style), Paragraph("<b>Status</b>", p_style)]
    ]
    for t_chk, t_proc, t_exp, t_act, t_stat in acceptance_tests:
        pdf_test_data.append([
            Paragraph(f"<b>{t_chk}</b>", p_style),
            Paragraph(t_exp, p_style),
            Paragraph(t_act, p_style),
            Paragraph(f"<b>{t_stat}</b>", p_style)
        ])
    pdf_test_tbl = Table(pdf_test_data, colWidths=[100, 140, 200, 60])
    pdf_test_tbl.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0F172A')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(pdf_test_tbl)
    story.append(Spacer(1, 10))

    # Section 5
    story.append(Paragraph("5. Conclusion & Technical Insights", h1_style))
    story.append(Paragraph("Completing Assignment #6 deepened my technical mastery of DOM manipulation, coordinate measurement, and positional animation sequencing using jQuery under strict zero-CSS-transition constraints. Managing live mutable application state and synchronizing data arrays with the DOM established robust engineering foundations for complex web interfaces.", p_style))

    pdf_doc.build(story)
    print(f"PDF exported successfully via ReportLab at: {pdf_path}")
except Exception as e:
    print(f"ReportLab PDF export failed: {e}")

